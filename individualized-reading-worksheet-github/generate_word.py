from __future__ import annotations

import argparse
import json
import os
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

PAGE_W = Cm(21)
PAGE_H = Cm(29.7)
MARGIN = Cm(1.45)
BODY_FONT = 'Microsoft JhengHei'
TITLE_FONT = 'Microsoft JhengHei'
KEYWORD_GREY = RGBColor(170, 170, 170)
ANSWER_RED = RGBColor(180, 30, 30)
DARK = RGBColor(35, 35, 35)


def set_cell_shading(cell, fill: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tcMar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader')
    tblHeader.set(qn('w:val'), 'true')
    trPr.append(tblHeader)


def set_run_font(run, name=BODY_FONT, size=Pt(13), bold=False, color=DARK):
    run.font.name = name
    run.font.size = size
    run.font.bold = bold
    run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.rFonts
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    for attr in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        rFonts.set(qn(f'w:{attr}'), name)


def style_paragraph(p, after=Pt(5), before=Pt(0), line=1.25):
    p.paragraph_format.space_after = after
    p.paragraph_format.space_before = before
    p.paragraph_format.line_spacing = line


def add_text(doc, text, size=13, bold=False, color=DARK, align=None, after=4):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    style_paragraph(p, after=Pt(after))
    r = p.add_run(text)
    set_run_font(r, size=Pt(size), bold=bold, color=color)
    return p


def add_heading_bar(doc, text, level=1):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    cell = table.cell(0, 0)
    set_cell_shading(cell, 'EEF4F1')
    set_cell_margins(cell, top=80, bottom=80)
    p = cell.paragraphs[0]
    style_paragraph(p, after=Pt(0))
    r = p.add_run(text)
    set_run_font(r, size=Pt(15 if level == 1 else 13), bold=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_image(doc, image_path: str | None, width_cm=14.5):
    if not image_path:
        return
    pth = Path(image_path)
    if not pth.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    r.add_picture(str(pth), width=Cm(width_cm))
    p.paragraph_format.space_after = Pt(5)


def add_keywords(doc, keywords: list[str] | None):
    if not keywords:
        return
    p = doc.add_paragraph()
    style_paragraph(p, after=Pt(7))
    r = p.add_run('關鍵語詞：' + '、'.join(keywords))
    set_run_font(r, size=Pt(11.5), color=KEYWORD_GREY)


def add_choices(doc, choices: list[str], answer: int | None = None, teacher=False, option_prefix='○'):
    for i, ch in enumerate(choices):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.45)
        style_paragraph(p, after=Pt(2), line=1.15)
        prefix = option_prefix
        r = p.add_run(f'{prefix} {ch}')
        is_answer = teacher and answer is not None and i == answer
        set_run_font(r, size=Pt(12.5), bold=is_answer, color=ANSWER_RED if is_answer else DARK)
        if is_answer:
            ra = p.add_run('  ✓')
            set_run_font(ra, size=Pt(12.5), bold=True, color=ANSWER_RED)


def add_question(doc, q: dict[str, Any], teacher=False):
    p = doc.add_paragraph()
    style_paragraph(p, after=Pt(3))
    r = p.add_run(q.get('prompt', ''))
    set_run_font(r, size=Pt(13), bold=True)
    qtype = q.get('type', 'choice')
    if qtype == 'choice':
        add_choices(doc, q.get('choices', []), q.get('answer'), teacher)
    elif qtype == 'fill':
        p2 = doc.add_paragraph()
        style_paragraph(p2, after=Pt(3))
        txt = q.get('line', '________________________________')
        r2 = p2.add_run(txt)
        set_run_font(r2, size=Pt(13))
        if q.get('word_bank'):
            rk = p2.add_run('    詞庫：' + '／'.join(q['word_bank']))
            set_run_font(rk, size=Pt(11.5), color=KEYWORD_GREY)
        if teacher and q.get('answer_text'):
            ans = doc.add_paragraph()
            rr = ans.add_run('答案：' + q['answer_text'])
            set_run_font(rr, size=Pt(11.5), bold=True, color=ANSWER_RED)
    elif qtype == 'short':
        lines = int(q.get('lines', 2))
        for _ in range(lines):
            p2 = doc.add_paragraph('_______________________________________________')
            style_paragraph(p2, after=Pt(4))
            for r2 in p2.runs:
                set_run_font(r2, size=Pt(13))
        if teacher and q.get('answer_text'):
            ans = doc.add_paragraph()
            rr = ans.add_run('參考答案：' + q['answer_text'])
            set_run_font(rr, size=Pt(11.5), bold=True, color=ANSWER_RED)


def add_table_block(doc, block: dict[str, Any], teacher=False):
    rows = block.get('rows', [])
    headers = block.get('headers', [])
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, 'EEF4F1')
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        set_run_font(r, size=Pt(11.5), bold=True)
        set_cell_margins(cell)
    answers = block.get('answers', [])
    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        for ci, val in enumerate(row):
            cell = cells[ci]
            set_cell_margins(cell, top=120, bottom=120)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            set_run_font(r, size=Pt(11.5))
            if teacher and ri < len(answers) and ci < len(answers[ri]) and answers[ri][ci]:
                p2 = cell.add_paragraph()
                rr = p2.add_run(answers[ri][ci])
                set_run_font(rr, size=Pt(11.5), bold=True, color=ANSWER_RED)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_3d_vocab(doc, vocab: dict[str, Any], teacher=False):
    add_heading_bar(doc, f"核心語詞｜{vocab.get('word','')}", level=2)
    add_image(doc, vocab.get('image'), width_cm=7.8)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(vocab.get('word',''))
    set_run_font(r, size=Pt(24), bold=True)
    if vocab.get('bopomofo'):
        r2 = p.add_run('   ' + vocab['bopomofo'])
        set_run_font(r2, size=Pt(14))
    add_text(doc, '生活例句：' + vocab.get('example',''), size=12.5)
    for ex in vocab.get('exercises', []):
        add_question(doc, ex, teacher=teacher)


def add_cover(doc, data, teacher=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(p, after=Pt(5))
    r = p.add_run(data.get('title', '閱讀理解學習單'))
    set_run_font(r, name=TITLE_FONT, size=Pt(22), bold=True)
    subtitle = f"{data.get('level','')} {'教師解答版' if teacher else '學生版'}"
    add_text(doc, subtitle, size=13, bold=True, color=KEYWORD_GREY, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run('姓名：________________    日期：________________')
    set_run_font(r2, size=Pt(12.5))
    p2.paragraph_format.space_after = Pt(8)


def generate(data: dict[str, Any], output_path: str, teacher=False):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = PAGE_W
    sec.page_height = PAGE_H
    sec.top_margin = MARGIN
    sec.bottom_margin = MARGIN
    sec.left_margin = MARGIN
    sec.right_margin = MARGIN

    styles = doc.styles
    styles['Normal'].font.name = BODY_FONT
    styles['Normal'].font.size = Pt(13)

    add_cover(doc, data, teacher=teacher)

    intro = data.get('prediction')
    if intro:
        add_heading_bar(doc, '一、從課文名稱預測課文內容')
        if intro.get('image'):
            add_image(doc, intro.get('image'), width_cm=12.5)
        add_question(doc, intro, teacher=teacher)

    section_num = 2
    for idx, block in enumerate(data.get('sections', []), start=1):
        add_heading_bar(doc, f"{section_num}、念一念課文")
        section_num += 1
        add_image(doc, block.get('image'))
        for para in block.get('text', []):
            add_text(doc, para, size=13)
        add_keywords(doc, block.get('keywords'))

        if data.get('level','').startswith('Level 3D') and block.get('vocab'):
            add_3d_vocab(doc, block['vocab'], teacher=teacher)

        if block.get('questions'):
            add_heading_bar(doc, f"{section_num}、重點提問")
            section_num += 1
            for q in block['questions']:
                add_question(doc, q, teacher=teacher)

        if block.get('fills'):
            add_heading_bar(doc, f"{section_num}、重點內容填寫")
            section_num += 1
            for q in block['fills']:
                add_question(doc, q, teacher=teacher)

        if block.get('content_questions'):
            add_heading_bar(doc, f"{section_num}、課文內容")
            section_num += 1
            for q in block['content_questions']:
                add_question(doc, q, teacher=teacher)

    if data.get('summary_table'):
        add_heading_bar(doc, f"{section_num}、整理表格")
        section_num += 1
        add_table_block(doc, data['summary_table'], teacher=teacher)

    if data.get('challenge'):
        add_heading_bar(doc, f"{section_num}、課文內容小挑戰")
        for q in data['challenge']:
            add_question(doc, q, teacher=teacher)

    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    try:
        doc.save(out)
    except PermissionError:
        alt = out.with_name(out.stem + '_更新版' + out.suffix)
        doc.save(alt)
        out = alt
    return str(out)


def main():
    ap = argparse.ArgumentParser(description='Generate individualized reading worksheet DOCX from JSON spec.')
    ap.add_argument('input_json')
    ap.add_argument('--output', required=True)
    ap.add_argument('--teacher', action='store_true')
    ap.add_argument('--auto-images', action='store_true', help='Generate missing context images before building Word (requires OPENAI_API_KEY).')
    ap.add_argument('--image-dir', default='images/generated')
    ap.add_argument('--image-cache-dir', default='images/cache')
    ap.add_argument('--image-model', default='gpt-image-2')
    ap.add_argument('--image-quality', choices=['low','medium','high'], default='medium')
    ap.add_argument('--overwrite-images', action='store_true')
    args = ap.parse_args()
    data = json.loads(Path(args.input_json).read_text(encoding='utf-8'))

    if args.auto_images:
        from generate_images import OpenAIImageGenerator, cached_generate, plan_images
        if not os.getenv('OPENAI_API_KEY'):
            raise SystemExit('OPENAI_API_KEY is required when --auto-images is enabled.')
        generator = OpenAIImageGenerator(args.image_model)
        jobs = plan_images(data, Path(args.image_dir))
        for n, job in enumerate(jobs, 1):
            path = cached_generate(
                generator,
                job['prompt'],
                job['path'],
                cache_dir=Path(args.image_cache_dir),
                size=job['size'],
                quality=args.image_quality,
                overwrite=args.overwrite_images,
            )
            target, key = job['target']
            target[key] = str(path.resolve())
            print(f'[image {n}/{len(jobs)}] {path}')

    print(generate(data, args.output, teacher=args.teacher))

if __name__ == '__main__':
    main()
